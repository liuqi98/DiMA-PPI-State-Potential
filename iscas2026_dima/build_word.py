"""Build an editable, single-column Word manuscript from the current LaTeX."""
from pathlib import Path
import re
import tempfile
import pypandoc
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
main = (ROOT/'dima_iscas2026_full.tex').read_text()
bbl = (ROOT/'dima_iscas2026_full.bbl').read_text()
entries = re.split(r'\\bibitem\{([^}]+)\}', bbl)[1:]
refs = {entries[i]: (i//2+1, entries[i+1].split(r'\end{thebibliography}')[0]) for i in range(0,len(entries),2)}
labels = dict(re.findall(r'\\newlabel\{([^}]+)\}\{\{([^}]+)\}', (ROOT/'dima_iscas2026_full.aux').read_text()))
text = main.split(r'\begin{abstract}')[1].split(r'\bibliographystyle')[0]
text = r'\section*{Abstract}'+'\n'+text
text = text.replace(r'\end{abstract}', '')
text = text.replace(r'\begin{IEEEkeywords}', r'\paragraph{Keywords:}').replace(r'\end{IEEEkeywords}', '')
def expand(m):
    p = ROOT/(m[1]+'.tex')
    return re.sub(r'\\input\{([^}]+)\}', expand, p.read_text())
text = re.sub(r'\\input\{([^}]+)\}', expand, text)
text = re.sub(r'\\cite\{([^}]+)\}', lambda m: '['+', '.join(str(refs[k][0]) for k in m[1].split(','))+']', text)
text = re.sub(r'\\eqref\{([^}]+)\}', lambda m:'('+labels[m[1]]+')', text)
text = re.sub(r'\\ref\{([^}]+)\}', lambda m:labels[m[1]], text)
# Explicit captions survive conversion independently of LaTeX float handling.
fig_no = 0
def figure(m):
    global fig_no
    fig_no += 1
    s=m[1]
    asset=re.search(r'\\includegraphics\[[^]]*\]\{([^}]+)\}',s)[1]
    cap=s.split(r'\caption{',1)[1].rsplit(r'\label',1)[0].strip().removesuffix('}').strip()
    if asset == 'fig_model.pdf':
        asset = 'model figure.jpg'
        cap = ('DiMA-PPI architecture. (a) Four bound/free and WT/mutant states are evaluated '
               'by a shared network; differences of state scores yield the predicted mutation effect. '
               '(b) Gated fusion integrates sequence, backbone and scalar features. '
               '(c) Each graph block combines four-head TransformerConv, SiLU, a residual '
               'connection and LayerNorm. Three graph blocks are stacked with separate parameters. '
               'Site and context readouts are concatenated with projected atomic descriptors '
               'before the 288--96--1 scalar head. Canonical state scores are reused across edges.')
    return '\n\\begin{center}\n\\includegraphics[width=6.7in]{'+asset.replace('.pdf','.png')+'}\n\\end{center}\n\n\\textbf{Fig. '+str(fig_no)+'.} '+cap+'\n\n'
text=re.sub(r'\\begin\{figure\*\}\[t\](.*?)\\end\{figure\*\}',figure,text,flags=re.S)
# Parse the caption's balanced braces (math superscripts contain braces).
def extract_braced(s, start):
    depth=1;i=start
    while depth:
        if s[i]=='{':depth+=1
        elif s[i]=='}':depth-=1
        i+=1
    return s[start:i-1],i
nt=0
def table(m):
    global nt
    nt+=1;s=m[1];start=s.index(r'\caption{')+9
    cap,end=extract_braced(s,start)
    tab=s[s.index(r'\begin{tabular}'):]
    tab=tab.replace(r'\!','')
    return '\n\\textbf{Table '+['I','II'][nt-1]+'.} '+cap+'\n\n'+tab+'\n'
text=re.sub(r'\\begin\{table\}\[t\](.*?)\\end\{table\}',table,text,flags=re.S)
text=re.sub(r'\\label\{[^}]+\}', '',text)
# Number all display equations consistently with the original PDF.
eq=0
def equation(m):
    global eq
    env,body=m[1],m[2]
    if env=='equation':
        eq+=1
        return r'\['+body+r'\qquad\text{('+str(eq)+r')}\]'
    parts=body.split(chr(92)*2)
    lines=[]
    for part in parts:
        if not part.strip():continue
        if r'\nonumber' in part:part=part.replace(r'\nonumber','')
        else:
            eq+=1;part+=r'\qquad\text{('+str(eq)+r')}'
        lines.append(part)
    return '\n\n'.join(r'\['+line.replace('&','')+r'\]' for line in lines)
text=re.sub(r'\\begin\{(equation|align)\}(.*?)\\end\{\1\}',equation,text,flags=re.S)
assert eq==9,eq
text=re.sub(r'\{\\rm\s+([^}]+)\}', lambda m: r'\mathrm{'+m[1]+'}', text)
text=text.replace(r'\operatorname{LN}', r'\text{LN}')
text=text.replace(r'|\widehat y_{a\to b}+\widehat y_{b\to a}|', r'\left|\widehat y_{a\to b}+\widehat y_{b\to a}\right|')
text=text.replace(r'|\widehat y_{a\to c}-\widehat y_{a\to b}-\widehat y_{b\to c}|', r'\left|\widehat y_{a\to c}-\widehat y_{a\to b}-\widehat y_{b\to c}\right|')
text+='\n\\section*{References}\n'
for key,(num,ref) in refs.items():
    ref=re.sub(r'\\BIBentry(?:ALT|STD)interwordspacing','',ref).strip()
    text+='\n['+str(num)+'] '+ref+'\n'
text=re.sub(r'(?m)^%.*\n','',text)
header=r'''\begin{center}
{\LARGE DiMA-PPI: Canonical State-Potential Learning for Consistent Protein Binding Mutation Prediction}

Qi Liu$^{1}$, Xinyi Xu$^{1}$, Qian Cheng$^{1}$, Chen Wu$^{1}$, Lu Meng$^{1}$,
Defu Yang$^{2}$, Xinbo Gao$^{1}$, and Xueli Chen$^{1,2}$

$^{1}$Xidian University, Xi'an, China

$^{2}$Hangzhou Normal University, Hangzhou, China

Correspondence: xyxu@xidian.edu.cn; xbgao@mail.xidian.edu.cn; xlchen@xidian.edu.cn
\end{center}
'''
text=re.sub(r'\$\^\{([^}]+)\}\$', lambda m: r'\textsuperscript{'+m[1]+'}', text)
header=re.sub(r'\$\^\{([^}]+)\}\$', lambda m: r'\textsuperscript{'+m[1]+'}', header)
with tempfile.TemporaryDirectory(prefix='dima_word_') as tmp:
    tex=Path(tmp)/'word.tex';tex.write_text(header+text)
    dest=ROOT/'dima_iscas2026_full.docx'
    pypandoc.convert_file(str(tex),'docx',format='latex',outputfile=str(dest),extra_args=['--standalone','--resource-path='+str(ROOT)])

doc=Document(dest)
for sec in doc.sections:
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=Inches(.75);sec.bottom_margin=Inches(.75)
    sec.left_margin=Inches(.85);sec.right_margin=Inches(.85)
for style in doc.styles:
    if style.type==1:
        style.font.name='Times New Roman';style.font.size=Pt(11)
        style.font.color.rgb=RGBColor(0,0,0)
        fonts=style.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in list(fonts.attrib):
            if 'Theme' in attr: del fonts.attrib[attr]
        fonts.set(qn('w:ascii'),'Times New Roman');fonts.set(qn('w:hAnsi'),'Times New Roman')
normal=doc.styles['Normal'];normal.paragraph_format.space_after=Pt(5)
normal.paragraph_format.line_spacing=1.05
for n in ['Heading 1','Heading 2']:
    doc.styles[n].font.size=Pt(12 if n=='Heading 1' else 11)
    doc.styles[n].font.bold=True
    doc.styles[n].paragraph_format.space_before=Pt(10)
    doc.styles[n].paragraph_format.space_after=Pt(4)
    doc.styles[n].paragraph_format.keep_with_next=True
for i,p in enumerate(doc.paragraphs):
    if i<5:p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    elif p.style.name=='Normal':p.alignment=WD_ALIGN_PARAGRAPH.JUSTIFY
    if i==0:
        for run in p.runs:run.font.size=Pt(17);run.bold=True
    if p.text.startswith('Fig. '):
        p.paragraph_format.space_after=Pt(9)
        for run in p.runs:run.font.size=Pt(9)
    if p.text.startswith('Table '):p.paragraph_format.keep_with_next=True
    if re.match(r'^\[\d+\]',p.text):
        for run in p.runs:run.font.size=Pt(9)
        p.paragraph_format.space_after=Pt(4)
    if p._p.xpath('.//w:drawing'):p.paragraph_format.keep_with_next=True
    if p.text.rstrip().endswith(('as','by','is',':')):p.paragraph_format.keep_with_next=True
for table in doc.tables:
    table.autofit=False
    for row in table.rows:
        trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after=Pt(3);p.paragraph_format.space_before=Pt(3)
                for run in p.runs:run.font.size=Pt(10)
    header=OxmlElement('w:tblHeader');table.rows[0]._tr.get_or_add_trPr().append(header)
footer=doc.sections[0].footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
doc.core_properties.title='DiMA-PPI: Canonical State-Potential Learning for Consistent Protein Binding Mutation Prediction'
doc.core_properties.author='Qi Liu et al.'
doc.save(dest)
assert len(doc.tables)==2
assert len(doc.inline_shapes)==2
print(dest)
print('9 numbered equations; 2 editable tables; 2 figures; 12 references.')
